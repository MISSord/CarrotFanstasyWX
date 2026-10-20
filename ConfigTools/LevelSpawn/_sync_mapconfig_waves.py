#!/usr/bin/env python3
"""将 Assets/Game/Json/Level 下 roundInfo 数量同步到 TbMapLevel.TotalRound（MapLevelInfo.xlsx / MapLevel）。"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[2]
LEVEL_DIR = ROOT / "CarrotFantasy" / "Assets" / "Game" / "Json" / "Level"
XLSX_PATH = ROOT / "ConfigTools" / "Config" / "Datas" / "MapLevelInfo.xlsx"
SHEET_NAME = "MapLevel"


def collect_wave_counts():
    counts = {}
    if not LEVEL_DIR.exists():
        raise FileNotFoundError("关卡目录不存在: " + str(LEVEL_DIR))
    for path in sorted(LEVEL_DIR.glob("Level*.json")):
        m = re.match(r"Level(\d+)_(\d+)\.json", path.name, re.I)
        if not m:
            continue
        big, lvl = int(m.group(1)), int(m.group(2))
        text = path.read_text(encoding="utf-8-sig")
        data = json.loads(text)
        counts[(big, lvl)] = len(data.get("roundInfo") or [])
    return counts


def find_header_map(ws):
    headers = {}
    for col in range(1, ws.max_column + 1):
        name = ws.cell(1, col).value
        if isinstance(name, str) and name:
            headers[name] = col
    return headers


def sync_xlsx(wave_counts: dict) -> int:
    if not XLSX_PATH.exists():
        raise FileNotFoundError("缺少 " + str(XLSX_PATH))

    wb = openpyxl.load_workbook(XLSX_PATH)
    if SHEET_NAME not in wb.sheetnames:
        raise RuntimeError("MapLevelInfo.xlsx 缺少子表 " + SHEET_NAME)
    ws = wb[SHEET_NAME]
    headers = find_header_map(ws)
    for required in ("BigLevelId", "LevelId", "TotalRound"):
        if required not in headers:
            raise RuntimeError("MapLevelInfo.xlsx/" + SHEET_NAME + " 缺少列 " + required)

    col_big = headers["BigLevelId"]
    col_level = headers["LevelId"]
    col_round = headers["TotalRound"]
    changed = 0
    seen = set()

    for row in range(6, ws.max_row + 1):
        big = ws.cell(row, col_big).value
        level = ws.cell(row, col_level).value
        if big is None or level is None:
            continue
        key = (int(big), int(level))
        seen.add(key)
        if key not in wave_counts:
            continue
        new_count = wave_counts[key]
        old = ws.cell(row, col_round).value
        if old != new_count:
            print(f"MapLevelInfo/{SHEET_NAME}: {key[0]}-{key[1]} {old} -> {new_count}")
            ws.cell(row, col_round).value = new_count
            changed += 1

    missing = sorted(k for k in wave_counts if k not in seen)
    for big, level in missing:
        print(f"warning: Level JSON {big}-{level} 在 MapLevelInfo/{SHEET_NAME} 中没有对应行")

    if changed:
        wb.save(XLSX_PATH)
    wb.close()
    print(f"{XLSX_PATH}: updated {changed}")
    return changed


def main():
    wave_counts = collect_wave_counts()
    total = sync_xlsx(wave_counts)
    print(f"done, total changes: {total}")
    print("请运行 ConfigTools/gen_code_json.bat 重新导出 Luban（编辑器菜单会自动导出）")


if __name__ == "__main__":
    main()
